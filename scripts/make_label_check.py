"""Build a blind labeling sheet to check the topic -> industry labels of BERTopic (Stage 3b) against K-Means (Stage 3).

200 random articles that both models assign to a sector. For each article the sheet shows the two labels as A and B
in random order; the key (which model is A) is written to a separate CSV. Score with scripts/score_label_check.py.
Both files stay under data/ because they contain titles and URLs from the corpus.
"""
from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from bertopic_topics import clean_title  # noqa: E402
from map_industries import MAP, NOISE  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data/labeling"
N, SEED = 200, 42
JUDGE = ["2 correct & specific", "1 correct but too general", "0 wrong"]
BETTER = ["A", "B", "tie"]
FONT, BOLD = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
INPUT = PatternFill("solid", start_color="FFFF00")


def main() -> None:
    d = pd.read_parquet(ROOT / "results/bertopic_v2/doc_topics_with_original_labels.parquet")
    d["sector"] = d["bertopic_topic_40"].map(lambda t: MAP[t][0] if t in MAP else None)
    pool = d[d.sector.notna() & (d.sector != NOISE) & (d.industry != "Noise")]
    s = pool.sample(n=N, random_state=SEED).reset_index(drop=True)

    # snippet = start of the Stage 2 cleaned article, not the AI-relevant sentences BERTopic saw, to stay neutral
    text = pd.read_parquet(ROOT / "data/clean_news_final_project.parquet", columns=["url", "text_clean"])
    text = text[text.url.isin(set(s.url))].drop_duplicates("url").set_index("url")["text_clean"]
    s["snippet"] = s.url.map(text).fillna("").str.slice(0, 500)
    s["title"] = s.title.map(clean_title)
    # K-Means labels use " / ", BERTopic labels use " & "; unify so the style does not reveal the model
    s["kmeans"] = s.industry.str.replace(" / ", " & ", regex=False)

    rng = np.random.default_rng(SEED)
    s["a_is"] = np.where(rng.random(N) < 0.5, "bertopic", "kmeans")
    s["label_A"] = np.where(s.a_is == "bertopic", s.sector, s.kmeans)
    s["label_B"] = np.where(s.a_is == "bertopic", s.kmeans, s.sector)
    s["id"] = np.arange(1, N + 1)

    OUT.mkdir(parents=True, exist_ok=True)
    s[["id", "url", "a_is", "sector", "industry"]].rename(
        columns={"sector": "bertopic_sector", "industry": "kmeans_industry"}).to_csv(OUT / "label_check_key.csv", index=False)

    wb = Workbook()
    ins = wb.active
    ins.title = "Instructions"
    lines = [
        ("Industry label check", BOLD),
        ("Each row is one news article with two candidate industry labels, A and B. Judge each label on its own, then pick the better one.", FONT),
        ("Fill in only the yellow columns on the Label sheet (H, I, J; K is optional notes). Use the dropdowns.", FONT),
        ("", FONT),
        ("Scale for A_judgment and B_judgment", BOLD),
        ("2 correct & specific  = the label names the article's main industry or sector.", FONT),
        ("1 correct but too general = not wrong, but would fit almost any AI article (e.g. a broad 'AI platforms' label for a school story).", FONT),
        ("0 wrong = the article is mainly about a different industry.", FONT),
        ("better = A, B, or tie (tie if both are equally good or equally bad).", FONT),
        ("", FONT),
        ("Example row (not on the Label sheet, for format only)", BOLD),
        ("title: 'Schools roll out AI tutors for maths homework' | label_A: Software & AI Platforms | label_B: Education", FONT),
        ("A_judgment: 1 correct but too general | B_judgment: 2 correct & specific | better: B", FONT),
        ("", FONT),
        ("Judge from the title and snippet; open the URL only if they are not enough. About 15-20 seconds per row.", FONT),
        ("Progress (rows with all three answers):", BOLD),
    ]
    for i, (text, font) in enumerate(lines, start=1):
        ins.cell(row=i, column=1, value=text).font = font
    ins.cell(row=len(lines), column=2, value=f'=SUMPRODUCT((Label!H2:H{N + 1}<>"")*(Label!I2:I{N + 1}<>"")*(Label!J2:J{N + 1}<>""))').font = BOLD
    ins.cell(row=len(lines), column=3, value=f"of {N}").font = BOLD
    ins.column_dimensions["A"].width = 130

    ws = wb.create_sheet("Label")
    cols = ["id", "title", "domain", "snippet", "url", "label_A", "label_B", "A_judgment", "B_judgment", "better", "notes"]
    widths = [5, 45, 18, 70, 30, 26, 26, 24, 24, 9, 25]
    for j, (c, w) in enumerate(zip(cols, widths), start=1):
        cell = ws.cell(row=1, column=j, value=c)
        cell.font = BOLD
        ws.column_dimensions[cell.column_letter].width = w
    for i, r in enumerate(s.itertuples(index=False), start=2):
        for j, v in enumerate([r.id, r.title, r.domain, r.snippet, r.url, r.label_A, r.label_B], start=1):
            cell = ws.cell(row=i, column=j, value=v)
            cell.font = FONT
            cell.alignment = Alignment(wrap_text=j in (2, 4), vertical="top")
        for j in (8, 9, 10, 11):
            cell = ws.cell(row=i, column=j)
            cell.fill = INPUT
            cell.font = FONT
            cell.alignment = Alignment(vertical="top")
    for letters, options in (("H2:I", JUDGE), ("J2:J", BETTER)):
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        dv.add(f"{letters}{N + 1}")
        ws.add_data_validation(dv)
    ws.freeze_panes = "B2"
    wb.calculation.fullCalcOnLoad = True  # the progress cell has no cached value until Excel recalculates
    wb.save(OUT / "industry_label_check.xlsx")
    print(f"wrote {OUT / 'industry_label_check.xlsx'} ({N} rows); A is BERTopic in {int((s.a_is == 'bertopic').sum())} rows")
    print(s.sector.value_counts().head(8).to_string()); print(s.industry.value_counts().to_string())


if __name__ == "__main__":
    main()
